# -*- coding: utf-8 -*-
"""通勤費の未申請者の抽出（経費チェックモードの独立ステップ ③）

②承認前精査の「通勤費申請なし」シートと同じ判定（kotsuhi_seisa.build_no_commute_rows）を、
**単独の成果物**として出す。②は承認が進むたびに何度も回す精査で、こちらは精査が終わった
（進行中0）あとに1回回して「当月に通勤費の申請が無い人」の一覧を確定させる。
判定ロジックは②と共用し、ここでは複製しない（2026-09-07 谷津さん依頼: 操作として分けたい）。

入力: ②と同じ（交通費申請CSV・①で出した経費チェックのブック）
出力: 通勤費申請なし_YYYY年M月.xlsx
      先頭シート＝一覧（メール下書きモードは先頭シートしか読まないので一覧を先頭に置く）
      2枚目＝実行情報（対象月・承認状況・使ったファイル・件数の内訳）
"""
from __future__ import annotations

from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font

from services.kotsuhi_seisa import (
    ACTIVE_STATUS,
    BORDER,
    HDR_FILL,
    HDR_FONT,
    _now_text,
    build_no_commute_rows,
    load_seisa_inputs,
    month_mismatch_warning,
    write_sheet,
)

SHEET_LIST = "通勤費申請なし"
SHEET_INFO = "実行情報"
# ②の同名シートから「前回比」だけ落とした列並び（こちらは1回で確定させる成果物なので差分は持たない）
LIST_COLUMNS = ["社員番号", "氏名", "出勤日数", "出社日数", "支給金額", "区分",
                "利用交通機関", "判定", "確認要否", "説明", "備考"]
JUDGE_ORDER = ("支給漏れの疑い", "実費申請の日数不一致", "勤怠実績なし", "マスタから支給")


@dataclass
class NoCommuteExtractResult:
    ok: bool
    output_path: Path
    month: str = ""
    error: "str | None" = None
    rows: list[dict] = field(default_factory=list)
    total: int = 0
    need_check: int = 0                      # 確認要否=要確認 の人数
    by_judge: dict = field(default_factory=dict)
    approved_rows: int = 0                   # 対象月の承認完了の明細行
    pending_rows: int = 0                    # 対象月の進行中（未承認）の明細行
    warnings: list[str] = field(default_factory=list)


def pending_warning(pending_rows: int) -> "str | None":
    """進行中の申請が残っているときの注意文。承認されると一覧から外れる人がいる。"""
    if pending_rows <= 0:
        return None
    return (f"進行中（未承認）の申請が {pending_rows} 行あります。承認が進むと一覧から外れる人が"
            f"いるため、②の精査が「進行中0」になってから抽出し直してください"
            f"（この一覧は現時点の申請状況で作っています）。")


def judge_counts(rows: list[dict]) -> dict:
    """判定ごとの人数。並びは JUDGE_ORDER（要確認が先）。"""
    c = Counter(r.get("判定", "") for r in rows)
    out = {k: c[k] for k in JUDGE_ORDER if c[k]}
    for k, v in c.items():
        if k not in out:
            out[k] = v
    return out


def write_no_commute_book(rows: list[dict], path: Path, info: dict) -> Path:
    """一覧（先頭）＋実行情報（2枚目）のブックを保存する。"""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    write_sheet(wb, SHEET_LIST, rows, LIST_COLUMNS,
                {"説明": 52, "判定": 16, "確認要否": 10}, severity_key="確認要否")
    ws = wb[SHEET_LIST]
    amt_col = LIST_COLUMNS.index("支給金額") + 1
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=amt_col).number_format = "#,##0"

    ws = wb.create_sheet(SHEET_INFO)
    ws.append(["項目", "内容"])
    for c in (1, 2):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font = HDR_FILL, HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
    for k, v in info.items():
        ws.append([k, v])
        for c in (1, 2):
            ws.cell(row=ws.max_row, column=c).border = BORDER
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 70

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def run_no_commute_extract(
    csv_path: "str | Path",
    check_xlsx: "str | Path",
    output_path: "str | Path",
    month: str,
    log_func=print,
    target_list: "Path | None" = None,
    excluded_list: "Path | None" = None,
) -> NoCommuteExtractResult:
    """当月に通勤費の申請が無い人を抽出し、単独のブックに書き出す。

    判定は②の build_no_commute_rows と同じ。②と違い前回比は付けず、実行のたびに
    同じファイル名で上書きする（精査が終わってから1回回す前提）。
    進行中の申請が残っているときは warnings に注意を入れるが、出力は作る。
    """
    csv_path, check_xlsx = Path(csv_path), Path(check_xlsx)
    output_path = Path(output_path)
    result = NoCommuteExtractResult(ok=False, output_path=output_path, month=month)

    if not csv_path.exists():
        result.error = f"交通費申請CSVが見つかりません: {csv_path}"
        return result
    if not check_xlsx.exists():
        result.error = f"経費チェックのブックが見つかりません: {check_xlsx}"
        return result
    try:
        y, m = month.split("-")
        month_text = f"{y}年{int(m)}月"
    except ValueError:
        result.error = f"対象月は YYYY-MM 形式で指定してください: {month}"
        return result

    try:
        src = load_seisa_inputs(csv_path, check_xlsx, month,
                                target_list=target_list, excluded_list=excluded_list)
        rows = build_no_commute_rows(src.details, src.idx, src.master, src.workdays,
                                     src.target_ids, src.excluded)
    except Exception as e:  # noqa: BLE001
        log_func(f"[error] 抽出に失敗しました: {e}")
        result.error = str(e)
        return result

    st = src.idx["ステータス"]
    result.approved_rows = sum(1 for r in src.details if r[st] == "承認完了")
    result.pending_rows = sum(1 for r in src.details if r[st] == "進行中")
    target_rows = sum(1 for r in src.details if r[st] in ACTIVE_STATUS)

    for w in (month_mismatch_warning(month, target_rows, src.out_of_month),
              pending_warning(result.pending_rows)):
        if w:
            result.warnings.append(w)
            log_func("[warn] " + w)

    result.rows = rows
    result.total = len(rows)
    result.by_judge = judge_counts(rows)
    result.need_check = sum(1 for r in rows if r.get("確認要否") == "要確認")

    info = {
        "実行日時": _now_text(),
        "対象月": month_text,
        "承認状況": f"承認完了 {result.approved_rows}行 / 進行中 {result.pending_rows}行",
        "交通費申請CSV": csv_path.name,
        "経費チェックブック": check_xlsx.name,
        "対象明細行(承認完了+進行中)": target_rows,
        "除外(対象月外の利用日)": src.out_of_month,
        "通勤費マスタ": f"{len(src.master.rows)}名 / {sum(len(v) for v in src.master.rows.values())}経路",
        "移動交通費（立替精算）対象者": f"{len(src.target_ids)}名" if src.target_ids else "なし",
        "精査対象外リスト": f"{len(src.excluded)}名" if src.excluded else "なし",
        "抽出人数": result.total,
        "うち要確認": result.need_check,
    }
    for k, v in result.by_judge.items():
        info[f"判定: {k}"] = f"{v}名"
    for i, w in enumerate(result.warnings, 1):
        info[f"注意{i}"] = w

    try:
        write_no_commute_book(rows, output_path, info)
    except Exception as e:  # noqa: BLE001
        log_func(f"[error] ブックの保存に失敗しました: {e}")
        result.error = str(e)
        return result

    result.ok = True
    log_func(f"[info] 承認状況: 承認完了 {result.approved_rows}行 / 進行中 {result.pending_rows}行")
    log_func(f"[info] 通勤費の未申請者: {result.total}名（うち要確認 {result.need_check}名）")
    for k, v in result.by_judge.items():
        log_func(f"[info]   {k}: {v}名")
    log_func(f"[done] 出力: {output_path}")
    return result
