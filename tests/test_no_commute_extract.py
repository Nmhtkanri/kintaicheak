# -*- coding: utf-8 -*-
"""通勤費の未申請者の抽出（独立ステップ③）のテスト。

判定そのものは kotsuhi_seisa.build_no_commute_rows のテストに任せ、ここでは
「単独の成果物として正しく出るか」「精査が終わっていないときに注意が出るか」を固定する。
"""
import csv
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import Workbook, load_workbook

from services.expense_check import _build_telework_sheets, add_commute_sheet
from services.no_commute_extract import (
    JUDGE_ORDER, SHEET_INFO, SHEET_LIST, judge_counts, pending_warning, run_no_commute_extract,
)

HEADER = ["ステータス", "交通機関", "社員番号", "申請者", "所属グループ", "申請書No.",
          "明細No.", "利用日", "金額", "往復", "小計", "乗車場所", "降車場所", "経路", "目的地"]


def _app_row(emp, kind="通勤定期代", status="承認完了", date="2026/07/01", amount="10000"):
    return [status, kind, emp, "申請者", "UAL", "1", "1", date, amount, "", amount, "A", "B", "", ""]


def _csv(tmp_path, rows):
    p = tmp_path / "交通費申請_9637_20260801.csv"
    with open(p, "w", encoding="cp932", newline="") as f:
        w = csv.writer(f)
        w.writerow(HEADER)
        w.writerows(rows)
    return p


def _commute(emp, name, amount, interval="毎月", no=1):
    return {"社員番号": emp, "氏名": name, "経路No": no, "出発": "自宅", "到着": "会社",
            "経由1": "", "経由2": "", "通勤経路": "自宅→会社", "利用交通機関": "電車",
            "支給間隔": interval, "支給方法": "", "支給金額": amount, "非課税通勤費": amount,
            "課税通勤費": 0, "支給開始": "2026-04-01", "片道距離(km)": ""}


def _book(tmp_path, summary, commute):
    wb = Workbook()
    _build_telework_sheets(wb, "2026-07", summary)
    add_commute_sheet(wb, commute)
    p = tmp_path / "経費チェック2026年7月.xlsx"
    wb.save(p)
    return p


def _emp(emp, name, work=20, tw=0):
    return {"id": emp, "name": name, "work_days": work,
            "telework_days": [(f"2026-07-{d:02d}", "テレワーク") for d in range(1, tw + 1)]}


def _run(tmp_path, app_rows, summary, commute, **kw):
    out = tmp_path / "out" / "通勤費申請なし_2026年7月.xlsx"
    return run_no_commute_extract(_csv(tmp_path, app_rows), _book(tmp_path, summary, commute),
                                  out, "2026-07", log_func=lambda *_: None, **kw)


def test_extract_writes_the_list_first_and_the_run_info_second(tmp_path):
    summary = [_emp("2020001", "定期 太郎"), _emp("2020002", "申請 花子"), _emp("2020003", "未登録 次郎")]
    commute = [_commute("2020001", "定期 太郎", 10000), _commute("2020002", "申請 花子", 8000)]
    res = _run(tmp_path, [_app_row("2020002")], summary, commute)

    assert res.ok, res.error
    assert res.output_path.exists()
    assert res.total == 2
    assert [r["社員番号"] for r in res.rows] == ["2020003", "2020001"]   # 要確認が先
    by = {r["社員番号"]: r for r in res.rows}
    assert by["2020001"]["判定"] == "マスタから支給"
    assert by["2020001"]["区分"] == "通勤定期代"
    assert by["2020001"]["支給金額"] == 10000
    assert by["2020003"]["判定"] == "支給漏れの疑い"
    assert by["2020003"]["確認要否"] == "要確認"
    assert res.need_check == 1
    assert res.by_judge == {"支給漏れの疑い": 1, "マスタから支給": 1}
    assert res.warnings == []

    wb = load_workbook(res.output_path)
    assert wb.sheetnames == [SHEET_LIST, SHEET_INFO]       # 一覧が先頭（メール下書きが先頭しか読まない）
    ws = wb[SHEET_LIST]
    header = [c.value for c in ws[1]]
    assert header[:2] == ["社員番号", "氏名"]
    assert "前回比" not in header
    assert ws.max_row == 3
    info = {r[0]: r[1] for r in wb[SHEET_INFO].iter_rows(min_row=2, values_only=True)}
    assert info["対象月"] == "2026年7月"
    assert info["抽出人数"] == 2
    assert info["うち要確認"] == 1
    assert info["承認状況"] == "承認完了 1行 / 進行中 0行"
    assert info["判定: マスタから支給"] == "1名"


def test_extract_warns_while_applications_are_still_pending(tmp_path):
    summary = [_emp("2020001", "定期 太郎"), _emp("2020002", "申請 花子")]
    commute = [_commute("2020001", "定期 太郎", 10000), _commute("2020002", "申請 花子", 8000)]
    res = _run(tmp_path, [_app_row("2020002", status="進行中")], summary, commute)

    assert res.ok, res.error
    assert res.pending_rows == 1 and res.approved_rows == 0
    assert any("進行中（未承認）の申請が 1 行" in w for w in res.warnings)
    # 出力は作る（注意付き）。実行情報にも注意を残す
    info = {r[0]: r[1] for r in load_workbook(res.output_path)[SHEET_INFO]
            .iter_rows(min_row=2, values_only=True)}
    assert "注意1" in info and "進行中" in info["注意1"]


def test_extract_keeps_travel_expense_members_as_such(tmp_path):
    members = tmp_path / "移動交通費対象者.csv"
    members.write_text("社員番号,氏名\n2020005,直行 五郎\n", encoding="utf-8-sig")
    res = _run(tmp_path, [], [_emp("2020005", "直行 五郎")], [], target_list=members)
    assert res.ok, res.error
    assert res.rows[0]["区分"] == "移動交通費"
    assert res.rows[0]["備考"] == "立替精算対象"


def test_extract_overwrites_the_same_file_without_diff(tmp_path):
    summary = [_emp("2020001", "定期 太郎")]
    commute = [_commute("2020001", "定期 太郎", 10000)]
    first = _run(tmp_path, [], summary, commute)
    second = _run(tmp_path, [], summary, commute)
    assert first.ok and second.ok
    assert first.output_path == second.output_path
    assert "前回比" not in second.rows[0]


def test_extract_reports_missing_inputs_in_japanese(tmp_path):
    res = run_no_commute_extract(tmp_path / "無い.csv", tmp_path / "無い.xlsx",
                                 tmp_path / "out.xlsx", "2026-07", log_func=lambda *_: None)
    assert not res.ok
    assert "交通費申請CSVが見つかりません" in res.error


def test_extract_explains_a_wrong_book(tmp_path):
    wb = Workbook()
    wb.active.title = "無関係"
    book = tmp_path / "別物.xlsx"
    wb.save(book)
    res = run_no_commute_extract(_csv(tmp_path, []), book, tmp_path / "out.xlsx", "2026-07",
                                 log_func=lambda *_: None)
    assert not res.ok
    assert "『通勤費』『サマリ』シートがありません" in res.error


def test_pending_warning_is_silent_at_zero():
    assert pending_warning(0) is None
    assert "3 行" in pending_warning(3)


def test_judge_counts_puts_need_check_first():
    rows = [{"判定": "マスタから支給"}, {"判定": "支給漏れの疑い"}, {"判定": "マスタから支給"},
            {"判定": "新しい判定"}]
    counts = judge_counts(rows)
    assert list(counts) == ["支給漏れの疑い", "マスタから支給", "新しい判定"]
    assert counts["マスタから支給"] == 2
    assert list(counts)[:2] == [k for k in JUDGE_ORDER if k in counts]
