# -*- coding: utf-8 -*-
"""kintai_import_runner の純粋関数テスト"""
from datetime import datetime

import pytest

from services.kintai_import_runner import (
    breaks_of_row,
    compare_row,
    is_kyuka_row,
    norm_date_iso,
    t2m,
    validate_upload_csv,
)


class TestT2m:
    def test_basic(self):
        assert t2m("8:45") == 525
        assert t2m("08:45") == 525

    def test_seconds_stripped(self):
        assert t2m("09:15:00") == 555

    def test_over_24h(self):
        """夜勤の24時超表記はそのまま分に換算する"""
        assert t2m("33:30") == 2010
        assert t2m("33:30:00") == 2010

    def test_empty_and_invalid(self):
        assert t2m("") is None
        assert t2m(None) is None
        assert t2m("1 day, 7:00:00") is None  # timedelta漏れは不正値として扱う


class TestNormDateIso:
    def test_formats(self):
        assert norm_date_iso("2026/6/1") == "2026-06-01"
        assert norm_date_iso("2026-06-01") == "2026-06-01"
        assert norm_date_iso("6/1/2026") == "2026-06-01"  # Excel米国式の巻き込み対策

    def test_invalid(self):
        assert norm_date_iso("") is None
        assert norm_date_iso("六月一日") is None


class TestIsKyukaRow:
    def test_no_kyuka(self):
        kyuka, why = is_kyuka_row({"休日休暇名1": "", "AM有休": "FALSE"})
        assert not kyuka

    def test_kyuka_name(self):
        kyuka, why = is_kyuka_row({"休日休暇名1": "年次有給"})
        assert kyuka
        assert "年次有給" in why

    def test_shubetsu_only(self):
        """守屋さん6/23パターン: 種別列にAM有休が入るケースを検知する"""
        kyuka, why = is_kyuka_row({"休日休暇名1": "年次有給", "休日休暇名1：種別": "AM有休"})
        assert kyuka
        assert "AM有休" in why

    def test_am_flag(self):
        kyuka, _ = is_kyuka_row({"AM有休": "TRUE"})
        assert kyuka

    def test_false_flags_ignored(self):
        kyuka, _ = is_kyuka_row({"AM有休": "FALSE", "PM有休": "0", "休日休暇名1": ""})
        assert not kyuka


class TestBreaksOfRow:
    def test_pair(self):
        row = {"休憩予定時刻1": "12:00", "復帰予定時刻1": "13:00"}
        assert breaks_of_row(row) == [(720, 780)]

    def test_half_pair_ignored(self):
        row = {"休憩予定時刻1": "12:00", "復帰予定時刻1": ""}
        assert breaks_of_row(row) == []

    def test_multiple_sorted(self):
        row = {"休憩予定時刻1": "15:00", "復帰予定時刻1": "15:30",
               "休憩予定時刻2": "12:00", "復帰予定時刻2": "13:00"}
        assert breaks_of_row(row) == [(720, 780), (900, 930)]

    def test_overnight_break(self):
        """夜勤の24時超休憩（30:00-31:30）"""
        row = {"休憩予定時刻1": "30:00", "復帰予定時刻1": "31:30"}
        assert breaks_of_row(row) == [(1800, 1890)]


class TestCompareRow:
    def test_all_match(self):
        row = {"出勤予定時刻": "8:50", "退勤予定時刻": "17:10",
               "休憩予定時刻1": "12:00", "復帰予定時刻1": "13:00",
               "出勤1": "8:50", "退勤1": "18:00"}
        sched = {"start": "08:50", "end": "17:10", "breaks": [("12:00", "13:00")]}
        att = {"in": "8:50", "out": "18:00"}
        assert compare_row(row, sched, att) == []

    def test_schedule_start_ng(self):
        row = {"出勤予定時刻": "8:50"}
        sched = {"start": "9:00", "breaks": []}
        ngs = compare_row(row, sched, {})
        assert len(ngs) == 1
        assert ngs[0][0] == "出勤予定"

    def test_empty_columns_not_compared(self):
        """CSVで空欄の項目は無処理なので比較しない"""
        row = {"出勤予定時刻": "", "退勤予定時刻": "", "出勤1": "", "退勤1": ""}
        assert compare_row(row, {"start": "9:00", "breaks": []}, {"in": "9:12"}) == []

    def test_punch_ng_has_marume_note(self):
        row = {"出勤1": "9:00"}
        ngs = compare_row(row, {}, {"in": "9:03"})
        assert len(ngs) == 1
        assert "まるめ前" in ngs[0][3]

    def test_breaks_ng(self):
        row = {"休憩予定時刻1": "12:00", "復帰予定時刻1": "13:00"}
        sched = {"start": "", "breaks": []}
        ngs = compare_row(row, sched, {})
        assert len(ngs) == 1
        assert ngs[0][0] == "休憩予定"
        assert "(なし)" in ngs[0][2]

    def test_missing_sched_and_att(self):
        row = {"出勤予定時刻": "8:50", "出勤1": "8:50"}
        ngs = compare_row(row, None, None)
        assert {n[0] for n in ngs} == {"出勤予定", "出勤打刻"}


class TestValidateUploadCsv:
    """送信前の門番チェック（2026-07-08 誤インポート事故の再発防止）"""

    TODAY = datetime(2026, 7, 13)

    def _rows(self, dates, extra=None):
        return [dict({"*年月日": d, "*従業員ID": "2011001", "出勤1": "9:00"},
                     **(extra or {})) for d in dates]

    def test_ok(self):
        errs = validate_upload_csv(self._rows(["2026/6/1", "2026/6/30"]),
                                   target_month="2026-06", month_explicit=True,
                                   today=self.TODAY)
        assert errs == []

    def test_us_date_blocked(self):
        """米国式(6/7/2026)はjinjerが日/月/年と誤解釈するため送信させない（事故の真因）"""
        errs = validate_upload_csv(self._rows(["6/7/2026"]), today=self.TODAY)
        assert any("米国式" in e for e in errs)

    def test_iso_date_blocked(self):
        errs = validate_upload_csv(self._rows(["2026-06-01"]), today=self.TODAY)
        assert any("YYYY/M/D" in e for e in errs)

    def test_seconds_cell_blocked(self):
        rows = self._rows(["2026/6/1"], extra={"退勤1": "33:30:00"})
        errs = validate_upload_csv(rows, target_month="2026-06",
                                   month_explicit=True, today=self.TODAY)
        assert any("秒付き" in e for e in errs)

    def test_multi_month_blocked(self):
        errs = validate_upload_csv(self._rows(["2026/6/30", "2026/7/1"]),
                                   today=self.TODAY)
        assert any("複数の月" in e for e in errs)

    def test_target_month_mismatch(self):
        errs = validate_upload_csv(self._rows(["2026/5/1"]),
                                   target_month="2026-06", month_explicit=True,
                                   today=self.TODAY)
        assert any("対象月" in e for e in errs)

    def test_future_month_blocked_even_if_explicit(self):
        errs = validate_upload_csv(self._rows(["2026/8/6"]),
                                   target_month="2026-08", month_explicit=True,
                                   today=self.TODAY)
        assert any("未来の月" in e for e in errs)

    def test_old_month_needs_explicit(self):
        errs = validate_upload_csv(self._rows(["2026/4/1"]), target_month="2026-04",
                                   month_explicit=False, today=self.TODAY)
        assert any("遡及" in e for e in errs)

    def test_old_month_allowed_when_explicit(self):
        errs = validate_upload_csv(self._rows(["2026/4/1"]), target_month="2026-04",
                                   month_explicit=True, today=self.TODAY)
        assert errs == []

    def test_prev_month_ok_without_explicit(self):
        """通常運用: 前月の締め修正は月指定なしで通る"""
        errs = validate_upload_csv(self._rows(["2026/6/15"]),
                                   target_month="2026-06", today=self.TODAY)
        assert errs == []


class TestStripScheduleForKyuka:
    """休暇日は予定を落として打刻だけ送る（2026-09-09 石橋9999999 AM有休で実測）。"""

    def test_schedule_and_template_blanked_punch_kept(self):
        from services.kintai_import_runner import strip_schedule_for_kyuka
        row = {"出勤予定時刻": "09:30", "退勤予定時刻": "17:30", "スケジュール雛形ID": "1",
               "休憩予定時刻1": "12:00", "復帰予定時刻1": "13:00",
               "出勤1": "09:30", "退勤1": "12:00", "休日休暇名1": "年次有給", "休日休暇名1：種別": "PM有休"}
        out, has_punch = strip_schedule_for_kyuka(row)
        assert has_punch
        assert out["出勤予定時刻"] == "" and out["退勤予定時刻"] == "" and out["スケジュール雛形ID"] == ""
        assert out["休憩予定時刻1"] == "" and out["復帰予定時刻1"] == ""
        assert out["出勤1"] == "09:30" and out["退勤1"] == "12:00"
        assert out["休日休暇名1"] == "年次有給"          # 休暇列は触らない
        assert row["出勤予定時刻"] == "09:30"            # 元の行は変更しない

    def test_no_punch_means_nothing_to_send(self):
        from services.kintai_import_runner import strip_schedule_for_kyuka
        row = {"出勤予定時刻": "09:30", "退勤予定時刻": "17:30", "出勤1": "", "退勤1": "  "}
        out, has_punch = strip_schedule_for_kyuka(row)
        assert not has_punch

    def test_break_only_counts_as_punch(self):
        from services.kintai_import_runner import strip_schedule_for_kyuka
        _out, has_punch = strip_schedule_for_kyuka({"休憩1": "12:00", "復帰1": "13:00"})
        assert has_punch


def test_dry_run_kyuka_rows_split(tmp_path):
    """dry-run: 休暇日で打刻ありは投入対象、打刻なしは対応不要（手動対応リストには載らない）。"""
    import csv
    from openpyxl import load_workbook
    from services.kintai_import_runner import run_api_import
    header = ["名前", "*従業員ID", "*年月日", "*打刻グループID", "スケジュール雛形ID",
              "出勤予定時刻", "退勤予定時刻", "出勤1", "退勤1", "休日休暇名1", "休日休暇名1：種別"]
    rows = [
        ["A", "1001", "2026/8/3", "40", "1", "09:30", "17:30", "09:30", "12:00", "年次有給", "PM有休"],
        ["A", "1001", "2026/8/4", "40", "1", "09:30", "17:30", "", "", "年次有給", "AM有休"],
        ["A", "1001", "2026/8/5", "40", "", "09:30", "17:30", "09:30", "18:00", "", ""],
    ]
    p = tmp_path / "up.csv"
    with open(p, "w", encoding="cp932", newline="") as f:
        w = csv.writer(f); w.writerow(header); w.writerows(rows)
    res = run_api_import(p, tmp_path, dry_run=True, month="2026-08", log_func=lambda m: None)
    assert res.ok
    assert res.submitted_rows == 2
    assert res.kyuka_punch_only == 1 and res.kyuka_skipped == 1
    assert res.excluded == []                                   # 手動対応リストは空
    wb = load_workbook(res.report_path, read_only=True)
    v = list(wb["検証結果"].iter_rows(values_only=True))[1:]
    assert any(r[2] == "2026-08-04" and r[3] == "対応不要" for r in v)
