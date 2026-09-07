# -*- coding: utf-8 -*-
"""経費チェック ③「通勤費の未申請者を抽出する」の画面配線とルート。

2026-09-07 谷津さん依頼: 未申請者の抽出は①の任意欄に混ぜず、精査（②）が終わったあとに
回す独立した項目にする。番号は ② → ③抽出 → ④仕訳追記 の並び。
"""
import csv
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import Workbook, load_workbook

import app as app_module
from config import Config
from services.expense_check import _build_telework_sheets, add_commute_sheet


def _html():
    return app_module.app.test_client().get("/").get_data(as_text=True)


def test_step_is_its_own_section_between_review_and_shiwake():
    html = _html()
    i_review = html.index("② 承認前の交通費精査")
    i_extract = html.index("③ 通勤費の未申請者を抽出する")
    i_shiwake = html.index("④ 仕訳データへの定期代追記")
    assert i_review < i_extract < i_shiwake
    assert "③ 仕訳データへの定期代追記" not in html   # 旧番号が残っていない


def test_section_is_wired_to_the_route_and_the_stepper():
    html = _html()
    for el in ("nc-run-btn", "nc-status", "nc-result-area", "nc-warn", "nc-table",
               "nc-download-link", "nc-console", "nc-error-area",
               "nc-cnt-total", "nc-cnt-check", "nc-cnt-master", "nc-cnt-noattend", "nc-cnt-pending"):
        assert 'id="' + el + '"' in html, el + " が無い"
    assert "fetch('/no_commute_extract'" in html
    # 入力は②の欄を使い回す（③に独自の入力欄は置かない）
    assert "document.getElementById('pr-csv').value" in html
    assert "'③未申請者'" in html
    assert "st.no_commute" in html


def test_route_rejects_missing_inputs_in_japanese():
    res = app_module.app.test_client().post("/no_commute_extract", data={
        "month": "2026-07", "kotsuhi_csv": "", "check_xlsx": ""})
    body = res.get_json()
    assert res.status_code == 400 and not body["success"]
    assert any("②の欄" in e for e in body["errors"])


def test_route_writes_the_book_into_outputs(tmp_path, monkeypatch):
    monkeypatch.setattr(Config, "OUTPUT_FOLDER", str(tmp_path / "outputs"))
    # 共有フォルダの名簿は読めなくても止まらない（存在しないパスにしておく）
    monkeypatch.setattr(Config, "KEIHI_TRAVEL_EXPENSE_MEMBERS_CSV", str(tmp_path / "none1.csv"))
    monkeypatch.setattr(Config, "KOTSUHI_EXCLUDED_MEMBERS_CSV", str(tmp_path / "none2.csv"))

    csv_path = tmp_path / "交通費申請.csv"
    with open(csv_path, "w", encoding="cp932", newline="") as f:
        w = csv.writer(f)
        w.writerow(["ステータス", "交通機関", "社員番号", "申請者", "所属グループ", "申請書No.",
                    "明細No.", "利用日", "金額", "往復", "小計", "乗車場所", "降車場所", "経路", "目的地"])
        w.writerow(["進行中", "通勤定期代", "2020002", "申請 花子", "UAL", "1", "1",
                    "2026/07/01", "8000", "", "8000", "A", "B", "", ""])
    wb = Workbook()
    _build_telework_sheets(wb, "2026-07", [
        {"id": "2020001", "name": "定期 太郎", "work_days": 20, "telework_days": []},
        {"id": "2020002", "name": "申請 花子", "work_days": 20, "telework_days": []},
    ])
    add_commute_sheet(wb, [{
        "社員番号": "2020001", "氏名": "定期 太郎", "経路No": 1, "出発": "自宅", "到着": "会社",
        "経由1": "", "経由2": "", "通勤経路": "自宅→会社", "利用交通機関": "電車",
        "支給間隔": "毎月", "支給方法": "", "支給金額": 10000, "非課税通勤費": 10000,
        "課税通勤費": 0, "支給開始": "2026-04-01", "片道距離(km)": ""}])
    book = tmp_path / "経費チェック2026年7月.xlsx"
    wb.save(book)

    res = app_module.app.test_client().post("/no_commute_extract", data={
        "month": "2026-07", "kotsuhi_csv": str(csv_path), "check_xlsx": str(book)})
    body = res.get_json()
    assert res.status_code == 200, body
    assert body["success"]
    assert body["output_filename"] == "通勤費申請なし_2026年7月.xlsx"
    assert body["download_url"] == "/download/通勤費申請なし_2026年7月.xlsx"
    assert [r["社員番号"] for r in body["rows"]] == ["2020001"]
    assert body["stats"]["total"] == 1
    assert body["stats"]["pending_rows"] == 1
    assert any("進行中" in w for w in body["stats"]["warnings"])
    out = tmp_path / "outputs" / "通勤費申請なし_2026年7月.xlsx"
    assert out.exists()
    assert load_workbook(out).sheetnames[0] == "通勤費申請なし"
